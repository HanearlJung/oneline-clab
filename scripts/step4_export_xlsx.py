"""Phase 4 — 화면 데이터 + 원천 추출값을 엑셀 한 파일로 내보낸다.

공유용이므로 화면에 보이는 것(탭별 영업 대상)과
그 근거가 된 원문 수치(단위 그대로)를 함께 담는다.

산출: 법인영업_타깃발굴_데이터_{기준일}.xlsx
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent

NAVY = "101827"
HEAD_FILL = PatternFill("solid", start_color=NAVY)
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(bold=True, size=14, color=NAVY)
NOTE_FONT = Font(size=9, color="6B7280")
THIN = Side(style="thin", color="E5E9F0")
BORDER = Border(bottom=THIN)

EOK = '#,##0;(#,##0);"-"'        # 억원
RAW = '#,##0;(#,##0);"-"'        # 원문 단위 그대로
PCT = '0.0;(0.0);"-"'
QTY = '#,##0;(#,##0);"-"'


def sheet(wb, name, title, note=None):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    if note:
        ws["A2"] = note
        ws["A2"].font = NOTE_FONT
    ws.freeze_panes = "A5"
    return ws


def header(ws, cols, row=4):
    for i, (label, width, _fmt) in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=label)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 30


def write_rows(ws, cols, rows, start=5):
    for r, data in enumerate(rows, start):
        for i, ((_l, _w, fmt), v) in enumerate(zip(cols, data), 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BORDER
            if fmt:
                c.number_format = fmt
            if isinstance(v, str) and len(v) > 40:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{start + len(rows) - 1}"


# ------------------------------------------------------------------ 시트별

TAB_COLS = [
    ("기업명", 16, None), ("종목코드", 10, None), ("시장", 9, None), ("업종", 22, None),
    ("왜 대상인가", 40, None), ("상세 근거", 52, None),
    ("영업 규모(억원)", 15, EOK), ("표시 금액", 14, None),
    ("영업 상황", 14, None), ("기준 날짜", 12, None), ("시간 표기", 18, None),
    ("제안 방향", 42, None), ("기준일", 12, None),
    ("접수번호", 16, None), ("원문 링크", 46, None),
]


def tab_rows(tab):
    # 공유 문서이므로 영업 규모 큰 순으로 정렬해 둔다 (금액 없는 행은 뒤로)
    ordered = sorted(tab["rows"], key=lambda r: -(r.get("amount") or 0))
    for r in ordered:
        yield [r["name"], r.get("code"), r.get("market"), r.get("industry"),
               r["reason"], r.get("detail"), r.get("amount"), r.get("metricText"),
               r.get("situation"), r.get("dateValue"), r.get("timeText"),
               r.get("action"), r.get("asOf"), r.get("rcpNo"), r.get("sourceUrl")]


def main():
    site = json.loads((ROOT / "data" / "site_data.json").read_text(encoding="utf-8"))
    comps = json.loads((ROOT / "data" / "companies.json").read_text(encoding="utf-8"))["companies"]
    master = json.loads((ROOT / "data" / "market_master.json").read_text(encoding="utf-8"))
    recs = {p.stem: json.loads(p.read_text(encoding="utf-8"))
            for p in sorted((ROOT / "extracted").glob("*.json"))}
    base_date = site["meta"]["baseDate"]

    wb = Workbook()
    wb.remove(wb.active)

    # ---------------------------------------------------------------- 개요
    ws = sheet(wb, "개요", "법인영업 타깃 발굴 — 데이터 명세",
               "삼성증권 C-Lab PoC · 출처: 금융감독원 전자공시시스템(DART) 공개 공시")
    ws["A4"] = "기준일";      ws["B4"] = base_date
    ws["A5"] = "대상 기업";   ws["B5"] = site["meta"]["companyCount"]
    ws["A6"] = "생성 시각";   ws["B6"] = site["meta"]["generated"]
    ws["A7"] = "배포 주소";   ws["B7"] = "https://hanearljung.github.io/oneline-clab/"
    for r in range(4, 8):
        ws.cell(row=r, column=1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 46

    ws["A9"] = "탭별 커버리지"
    ws["A9"].font = Font(bold=True, size=11)
    cov_cols = [("탭", 18, None), ("해당 기업 수", 13, None), ("전체", 8, None),
                ("커버리지", 11, "0.0%"), ("행 수", 9, None), ("원천 공시", 40, None)]
    header(ws, cov_cols, row=10)

    SRC = {
        "investment": "정기보고서 주석 '금융상품·범주별 금융자산'",
        "pension": "정기보고서 주석 '종업원급여·확정급여제도'",
        "blockdeal": "주요사항보고서 '자기주식 처분·취득 결정'",
        "pledge": "주식등의 대량보유상황보고서 '보유주식등에 관한 계약'",
        "employee": "정기보고서 '주식매수선택권 부여 현황'",
        "wealth": "거래소공시 '최대주주변경' / '주식양수도 계약 체결'",
    }
    SHEETS = {"investment": "금융상품운용", "pension": "퇴직연금", "blockdeal": "자사주처분",
              "pledge": "주식담보대출", "employee": "임직원주식보상", "wealth": "대주주자금운용"}

    for i, t in enumerate(site["tabs"]):
        r = 11 + i
        nm = SHEETS[t["key"]]
        ws.cell(row=r, column=1, value=t["title"]).border = BORDER
        ws.cell(row=r, column=2, value=t["coverage"]["covered"]).border = BORDER
        ws.cell(row=r, column=3, value=t["coverage"]["total"]).border = BORDER
        c = ws.cell(row=r, column=4, value=f"=IF(C{r}=0,0,B{r}/C{r})")
        c.number_format = "0.0%"; c.border = BORDER
        c = ws.cell(row=r, column=5, value=f"=COUNTA({nm}!B5:B2000)")
        c.border = BORDER
        ws.cell(row=r, column=6, value=SRC[t["key"]]).border = BORDER

    r0 = 11 + len(site["tabs"])
    ws.cell(row=r0, column=1, value="합계 행 수").font = Font(bold=True)
    ws.cell(row=r0, column=5, value=f"=SUM(E11:E{r0 - 1})").font = Font(bold=True)

    ws.cell(row=r0 + 2, column=1, value="읽는 법").font = Font(bold=True, size=11)
    for i, line in enumerate([
        "· '영업 규모'는 회계상 총액이 아니라 증권사가 실제로 유치·전환할 수 있는 금액이다.",
        "  (원금보장형 적립금 / 예치성 자금 / 담보 대출금액 / 양수도 대금)",
        "· 값이 공시되지 않은 기업은 해당 탭에서 제외했다. 커버리지가 그 결과다.",
        "· 이미 지난 이벤트(처분 종료·담보 만기 경과)는 제외했다.",
        "· 기준일은 기업마다 다르다. 결산월과 공시 시점이 제각각이기 때문이다.",
        "· 원문 링크는 그 값을 뽑은 실제 공시로 연결된다.",
    ]):
        ws.cell(row=r0 + 3 + i, column=1, value=line).font = NOTE_FONT

    # ------------------------------------------------------------ 대상기업
    ws = sheet(wb, "대상기업", "대상 상장사 100개사",
               "이벤트 공시 보유 종류 우선, 동수면 시가총액 순으로 선정")
    cols = [("기업명", 18, None), ("종목코드", 10, None), ("고유번호", 11, None),
            ("시장", 9, None), ("업종", 26, None), ("시가총액(억원)", 15, EOK),
            ("최신 정기보고서", 24, None), ("접수일", 11, None),
            ("자사주", 8, None), ("담보", 8, None), ("최대주주변경", 12, None)]
    header(ws, cols)
    yn = lambda b: "O" if b else ""
    write_rows(ws, cols, [[
        c["name"], c.get("stock_code"), c["corp_code"], c["market"],
        c.get("industry"), c.get("market_cap"),
        c["report"]["report_nm"], c["report"]["rcept_dt"],
        yn(c["events"].get("treasury")), yn(c["events"].get("pledge")),
        yn(c["events"].get("owner_change")),
    ] for c in comps])

    # -------------------------------------------------------------- 탭 6종
    for t in site["tabs"]:
        nm = SHEETS[t["key"]]
        ws = sheet(wb, nm, f"{t['title']} — 영업 대상",
                   f"{t['desc']} · {t['coverage']['covered']}/{t['coverage']['total']}개사 해당 "
                   f"· 시간 열: {t['timeColumn']['label']}")
        header(ws, TAB_COLS)
        rows = list(tab_rows(t))
        if rows:
            write_rows(ws, TAB_COLS, rows)
        else:
            ws["A5"] = "해당 기업 없음 — 스톡옵션·RSU 를 운영하는 상장사 자체가 드물다."
            ws["A5"].font = NOTE_FONT

    # ------------------------------------------------------- 원천 추출값
    ws = sheet(wb, "원천_금융자산", "금융자산 — 원문 수치",
               "정기보고서 주석에서 직접 읽은 값. 단위는 원문 그대로.")
    cols = [("기업명", 16, None), ("단위", 8, None), ("금융자산 총계", 16, RAW),
            ("상각후원가", 16, RAW), ("FVOCI(전략지분)", 16, RAW), ("FVTPL(수익증권류)", 17, RAW),
            ("예치성 자금", 16, RAW), ("현금및현금성", 16, RAW),
            ("상품 구성", 34, None), ("보고서", 22, None), ("기준일", 11, None), ("접수번호", 16, None)]
    header(ws, cols)
    write_rows(ws, cols, [[
        r["name"], fa.get("unit"), fa.get("total"), fa.get("amortized_cost"),
        fa.get("fvoci"), fa.get("fvtpl"), fa.get("short_term_deposits"), fa.get("cash"),
        ", ".join(fa.get("products") or []),
        (r.get("source") or {}).get("report_nm"), (r.get("source") or {}).get("as_of"),
        (r.get("source") or {}).get("rcp_no"),
    ] for r in recs.values() if (fa := r.get("financial_assets"))])

    ws = sheet(wb, "원천_퇴직연금", "퇴직연금 — 원문 수치",
               "확정급여제도. 사외적립자산 구성이 곧 영업 세그먼트를 가른다.")
    cols = [("기업명", 16, None), ("단위", 8, None), ("확정급여채무(DBO)", 18, RAW),
            ("사외적립자산", 16, RAW), ("순확정급여부채", 16, RAW), ("적립률(%)", 11, PCT),
            ("원금보장형", 16, RAW), ("실적배당형", 16, RAW), ("기타", 13, RAW),
            ("당기 기여금", 15, RAW), ("차기 예상 기여금", 17, RAW),
            ("제도", 10, None), ("수탁기관", 20, None),
            ("보고서", 22, None), ("기준일", 11, None), ("접수번호", 16, None)]
    header(ws, cols)
    write_rows(ws, cols, [[
        r["name"], p.get("unit"), p.get("dbo"), p.get("plan_assets"),
        p.get("net_liability"), p.get("funding_ratio"),
        p.get("principal_guaranteed"), p.get("indirect_investment"), p.get("other_assets"),
        p.get("employer_contribution"), p.get("expected_contribution"),
        p.get("plan_type"), ", ".join(p.get("providers") or []),
        (r.get("source") or {}).get("report_nm"), (r.get("source") or {}).get("as_of"),
        (r.get("source") or {}).get("rcp_no"),
    ] for r in recs.values() if (p := r.get("pension"))])

    ws = sheet(wb, "원천_자사주", "자기주식 취득·처분 결정 — 원문 수치",
               "최근 12개월 결정 전건. 화면에는 예정일이 남은 건과 보유 자사주만 노출.")
    cols = [("기업명", 16, None), ("유형", 8, None), ("수량(주)", 14, QTY),
            ("금액(원)", 18, RAW), ("시작 예정일", 12, None), ("종료 예정일", 12, None),
            ("방법", 14, None), ("목적", 46, None), ("결정 전 보유 자사주(주)", 20, QTY),
            ("접수번호", 16, None)]
    header(ws, cols)
    write_rows(ws, cols, [[
        r["name"], e.get("event_type"), e.get("shares"), e.get("amount"),
        e.get("period_start"), e.get("period_end"), e.get("method"), e.get("purpose"),
        e.get("held_shares"), e.get("rcp_no"),
    ] for r in recs.values() for e in (r.get("treasury") or [])])

    ws = sheet(wb, "원천_주식담보", "대주주 주식담보 계약 — 원문 수치",
               "대량보유상황보고서 '보유주식등에 관한 계약'. 담보성 계약만 추출(양수도·공동보유 제외).")
    cols = [("기업명", 16, None), ("보유자", 18, None), ("관계", 14, None),
            ("계약 종류", 16, None), ("계약 상대방", 24, None),
            ("담보주식수(주)", 15, QTY), ("대출금액(원)", 18, RAW), ("지분율(%)", 11, PCT),
            ("계약체결일", 12, None), ("만기", 12, None), ("채무자", 18, None),
            ("접수번호", 16, None)]
    header(ws, cols)
    write_rows(ws, cols, [[
        r["name"], p.get("holder"), p.get("relation"), p.get("contract_type"),
        p.get("counterparty"), p.get("shares"), p.get("loan_amount"), p.get("ratio"),
        p.get("contract_date"), p.get("maturity"), p.get("debtor"), p.get("rcp_no"),
    ] for r in recs.values() for p in (r.get("pledge") or [])])

    ws = sheet(wb, "원천_최대주주변경", "최대주주 변경 — 원문 수치",
               "인수자금·양수도 대금이 곧 구 대주주에게 유입된 금액이다.")
    cols = [("기업명", 16, None), ("변경 전 최대주주", 26, None), ("변경 후 최대주주", 26, None),
            ("변경 전 지분(%)", 14, PCT), ("변경 후 지분(%)", 14, PCT),
            ("변경일", 12, None), ("거래금액(원)", 18, RAW),
            ("자기자금(원)", 18, RAW), ("차입금(원)", 18, RAW),
            ("사유", 46, None), ("접수번호", 16, None)]
    header(ws, cols)
    write_rows(ws, cols, [[
        r["name"], s.get("prev_largest"), s.get("largest"),
        s.get("prev_ratio"), s.get("largest_ratio"), s.get("change_date"),
        s.get("deal_amount"), s.get("deal_own"), s.get("deal_debt"),
        s.get("reason"), s.get("rcp_no"),
    ] for r in recs.values() if (s := r.get("shareholders")) and s.get("change_date")])

    out = ROOT / f"법인영업_타깃발굴_데이터_{base_date.replace('-', '')}.xlsx"
    wb.save(out)
    print(f"저장: {out.name}")
    for s in wb.sheetnames:
        print(f"  {s:16s} {wb[s].max_row - 4 if wb[s].max_row > 4 else 0:4d}행")


if __name__ == "__main__":
    main()
